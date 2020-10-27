"""
PrsAuditor is used to check the potential invalidity in the presentation (e.g. insufficient sample size),
it will scan the data_container, layout design structure, and other meta data to issue warnings.

For simplicity, the PrsAuditor pass the original mechanism to create the object on slide but add warning
on the slide directly. This workaround should be harmless as the auditor only create textbox on the slide.
"""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from utils.pptx_params import textbox


class PrsAuditor:
    def __init__(self):
        self.data = None
        self.log = {}

    def audit_data_container(self, data_container):
        audit = {}
        self.log['data_container'] = audit
        return audit

    def audit_design_structer(self, design_structure):
        audit = {}
        self.log['design_structure'] = audit
        return audit

    @staticmethod
    def _sample_size_warning(slide, left=Inches(8.33), top=Inches(1)):
        # activate when the auditor find chart with sample size less than pre-defined threshold

        # the position of the comment
        width = Inches(2.66)
        height = Inches(0.37)
        txBox = slide.shapes.add_textbox(left, top, width, height)

        # the color of the comment
        fill = txBox.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(200, 200, 200)

        # the content and the format of the comment
        tf = txBox.text_frame
        tf.clear()  # not necessary for newly-created shape
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = textbox("sample_warn_text")
        font = run.font
        font.name = textbox("sample_warn_font")
        font.size = Pt(16)
        return slide
